from __future__ import annotations

import base64
import csv
import io
import json
import os
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "data" / "stammdaten.json"
HOST = "127.0.0.1"
PORT = 8000

def load_stammdaten():
    if not DATA_FILE.exists():
        return {"version": 1, "linienbuendel": [], "verkehrsunternehmen": [], "linienzuordnungen": []}
    with DATA_FILE.open("r", encoding="utf-8") as f:
        return json.load(f)

def save_stammdaten(payload):
    if not isinstance(payload, dict):
        raise ValueError("Stammdaten muessen ein JSON-Objekt sein.")
    payload.setdefault("version", 1)
    payload.setdefault("linienbuendel", [])
    payload.setdefault("verkehrsunternehmen", [])
    payload.setdefault("linienzuordnungen", [])
    tmp = DATA_FILE.with_suffix(".json.tmp")
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    with tmp.open("w", encoding="utf-8", newline="\n") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
        f.write("\n")
    os.replace(tmp, DATA_FILE)


def _clean_cell(value):
    if value is None:
        return ""
    return str(value).strip()


def preview_import_file(filename, raw_bytes):
    suffix = Path(filename).suffix.lower()

    if suffix == ".csv":
        text = None
        encoding_used = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = raw_bytes.decode(encoding)
                encoding_used = encoding
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            raise ValueError("CSV-Zeichensatz konnte nicht gelesen werden.")

        sample = text[:20000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";"

        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        rows = [row for row in rows if any(str(cell).strip() for cell in row)]

        if not rows:
            raise ValueError("Die CSV-Datei enthält keine Daten.")

        headers = [_clean_cell(x) or f"Spalte {i+1}" for i, x in enumerate(rows[0])]
        data_rows = rows[1:]

        return {
            "format": f"CSV ({encoding_used}, Trennzeichen {repr(delimiter)})",
            "sheet_name": filename,
            "headers": headers,
            "preview_rows": [[_clean_cell(x) for x in row] for row in data_rows[:8]],
            "row_count": len(data_rows),
            "column_count": len(headers),
        }

    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError(
                "Excel-Dateien benötigen das Python-Paket openpyxl. "
                "CSV-Dateien können bereits geprüft werden."
            ) from exc

        wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]

        iterator = ws.iter_rows(values_only=True)
        first = next(iterator, None)
        if first is None:
            raise ValueError("Das erste Tabellenblatt enthält keine Daten.")

        headers = [_clean_cell(x) or f"Spalte {i+1}" for i, x in enumerate(first)]
        preview_rows = []
        row_count = 0

        for row in iterator:
            values = [_clean_cell(x) for x in row]
            if not any(values):
                continue
            row_count += 1
            if len(preview_rows) < 8:
                preview_rows.append(values)

        return {
            "format": "Excel",
            "sheet_name": ws.title,
            "headers": headers,
            "preview_rows": preview_rows,
            "row_count": row_count,
            "column_count": len(headers),
        }

    raise ValueError("Unterstützt werden derzeit CSV, XLSX und XLSM.")

class DashboardHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def _send_json(self, payload, status=HTTPStatus.OK):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _read_json_body(self):
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0:
            raise ValueError("Leerer Request.")
        raw = self.rfile.read(length)
        return json.loads(raw.decode("utf-8"))

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/api/health":
            self._send_json({"ok": True, "service": "AMINA Dashboard Server"})
            return
        if path == "/api/stammdaten":
            try:
                self._send_json(load_stammdaten())
            except Exception as exc:
                self._send_json({"ok": False, "error": str(exc)}, status=HTTPStatus.INTERNAL_SERVER_ERROR)
            return
        super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path

        if path == "/api/import/preview":
            try:
                payload = self._read_json_body()
                filename = str(payload.get("filename", "")).strip()
                encoded = payload.get("content_base64", "")

                if not filename or not encoded:
                    raise ValueError("Dateiname oder Dateiinhalt fehlt.")

                raw_bytes = base64.b64decode(encoded, validate=True)

                if len(raw_bytes) > 50 * 1024 * 1024:
                    raise ValueError("Die Datei ist größer als 50 MB.")

                preview = preview_import_file(filename, raw_bytes)
                preview.update({
                    "ok": True,
                    "filename": filename,
                    "linienbuendel_id": payload.get("linienbuendel_id", ""),
                    "verkehrsunternehmen_id": payload.get("verkehrsunternehmen_id", ""),
                    "datenart": payload.get("datenart", ""),
                    "liefermonat": payload.get("liefermonat", ""),
                })
                self._send_json(preview)
            except (ValueError, json.JSONDecodeError) as exc:
                self._send_json(
                    {"ok": False, "error": str(exc)},
                    status=HTTPStatus.BAD_REQUEST
                )
            except Exception as exc:
                self._send_json(
                    {"ok": False, "error": str(exc)},
                    status=HTTPStatus.INTERNAL_SERVER_ERROR
                )
            return

        if path == "/api/stammdaten":
            try:
                payload = self._read_json_body()
                save_stammdaten(payload)
                self._send_json({"ok": True, "saved": True})
            except (ValueError, json.JSONDecodeError) as exc:
                self._send_json({"ok": False, "error": str(exc)}, status=HTTPStatus.BAD_REQUEST)
            except Exception as exc:
                self._send_json({"ok": False, "error": str(exc)}, status=HTTPStatus.INTERNAL_SERVER_ERROR)
            return

        self._send_json(
            {"ok": False, "error": "Unbekannter API-Endpunkt."},
            status=HTTPStatus.NOT_FOUND
        )

    def end_headers(self):
        if self.path.startswith("/dashboard/"):
            self.send_header("Cache-Control", "no-store")
        super().end_headers()

if __name__ == "__main__":
    print(f"AMINA Dashboard Server: http://localhost:{PORT}/dashboard/")
    print(f"Stammdaten-API:        http://localhost:{PORT}/api/stammdaten")
    print("Zum Beenden: Strg+C")
    server = ThreadingHTTPServer((HOST, PORT), DashboardHandler)
    server.serve_forever()

