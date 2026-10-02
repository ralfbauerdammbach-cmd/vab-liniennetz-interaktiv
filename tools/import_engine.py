from __future__ import annotations

import csv
import io
import json
import re
import shutil
import uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from contract_rules import annotate_bundle
from typing import Any

try:
    import openpyxl
except ImportError:
    openpyxl = None


CANONICAL_FIELDS = [
    ("record_id", "Datensatz-ID"),
    ("trip_id", "Fahrtnummer / Fahrt-ID"),
    ("service_date", "Betriebstag / Datum"),
    ("day_type", "Tagesart"),
    ("operator", "Verkehrsunternehmen"),
    ("vehicle", "Fahrzeug / Kennzeichen"),
    ("line", "Linie"),
    ("direction", "Richtung"),
    ("stop_name", "Haltestelle"),
    ("stop_id", "Haltestellen-ID / DHID"),
    ("platform", "Steig / Bussteig"),
    ("scheduled_arrival", "Soll-Ankunft"),
    ("scheduled_departure", "Soll-Abfahrt"),
    ("actual_arrival", "Ist-Ankunft"),
    ("actual_departure", "Ist-Abfahrt"),
    ("distance", "Strecke"),
    ("journey_time", "Fahrzeit"),
    ("boarding", "Einsteiger"),
    ("alighting", "Aussteiger"),
    ("occupancy", "Besetzung"),
    ("load_factor", "Auslastung"),
    ("disruption", "Störhalt"),
    ("disposition", "Disposition / Änderungsart"),
]

ALIASES = {
    "record_id": ["id", "datensatz id", "zeilen id", "record id"],
    "trip_id": ["fahrt id", "fahrtid", "fahrtnummer", "fahrt nr", "fahrtnr", "journeyinfo", "servicejourneyid"],
    "service_date": ["fahrt datum", "fahrtdatum", "datum", "betriebstag", "service date", "servicedate", "timestamp"],
    "day_type": ["tagesart", "tagtyp", "daytype", "verkehrstagestyp"],
    "operator": ["verkehrsunternehmen", "operator", "operatorid", "vu"],
    "vehicle": ["kennzeichen", "fahrzeug", "fahrzeug id", "vehicle", "vehicleid"],
    "line": ["linie", "linien nr", "liniennummer", "linienkennung", "line", "lines", "dlid"],
    "direction": ["richtung", "fahrtrichtung", "direction"],
    "stop_name": ["haltestelle", "haltepunkt name", "haltestellenname", "stop", "stopname", "stopref"],
    "stop_id": ["dhid", "haltepunkt id", "haltestellen id", "stop id", "stopid"],
    "platform": ["steig", "bussteig", "mast", "platform"],
    "scheduled_arrival": ["ankunft soll", "soll ankunft", "sollankunft", "plan ankunft", "planankunftszeit"],
    "scheduled_departure": ["abfahrt soll", "soll abfahrt", "sollabfahrt", "plan abfahrt", "planabfahrtszeit"],
    "actual_arrival": ["ankunft ist", "ist ankunft", "istankunft", "tatsachliche ankunft"],
    "actual_departure": ["abfahrt ist", "ist abfahrt", "istabfahrt", "tatsachliche abfahrt"],
    "distance": ["strecke", "gefahrene strecke", "fahrt strecke", "distanz"],
    "journey_time": ["fahrzeit", "fahrt fahrzeit", "reisezeit"],
    "boarding": ["einsteiger", "einsteiger mittel", "einv", "einn"],
    "alighting": ["aussteiger", "aussteiger mittel", "ausv", "ausn"],
    "occupancy": ["belegung", "belegung mittel", "besetzung"],
    "load_factor": ["auslastung", "auslastungsgrad"],
    "disruption": ["storhalte", "störhalte", "stoerhalte", "störung", "stoerung"],
    "disposition": ["disposition", "anderungsart", "änderungsart", "massnahme", "maßnahme"],
}


def normalize_name(value: Any) -> str:
    s = str(value or "").strip().lower()
    table = str.maketrans({
        "ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss",
        "_": " ", "-": " ", "/": " ", ".": " ", ":": " ",
    })
    s = s.translate(table)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


NORMALIZED_ALIASES = {
    key: [normalize_name(x) for x in vals]
    for key, vals in ALIASES.items()
}



def _normalize_operator_token(value: Any) -> str:
    s = str(value or "").strip().lower()
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    s = re.sub(r"[^a-z0-9]+", "", s)
    return s


def line_derivation_rule(operator: dict[str, Any] | None) -> dict[str, Any] | None:
    """
    Fachregel:
    Nur für Ehrlich und DB-VU wird die Liniennummer aus den ersten zwei Stellen
    der Fahrt-ID abgeleitet.
    """
    if not operator:
        return None

    name = _normalize_operator_token(operator.get("name", ""))
    code = _normalize_operator_token(operator.get("kuerzel", ""))

    is_ehrlich = "ehrlich" in name or "ehrlich" in code
    is_db_vu = name in {"dbvu", "dbverkehrsunternehmen"} or code == "dbvu"

    if is_ehrlich or is_db_vu:
        return {
            "id": "trip_id_first_two_digits",
            "label": "Linie aus den ersten zwei Stellen der Fahrt-ID",
            "source_field": "trip_id",
            "digits": 2,
        }

    return None


def derive_line_values(
    rows: list[list[Any]],
    headers: list[str],
    mapping: dict[str, str],
    operator: dict[str, Any] | None,
) -> dict[str, Any]:
    rule = line_derivation_rule(operator)
    if not rule:
        return {
            "active": False,
            "rule": None,
            "values": [],
            "counter": Counter(),
            "source_header": None,
            "invalid_count": 0,
        }

    source_header = mapping.get("trip_id")
    source_idx = get_col_index(headers, source_header)
    if source_idx is None:
        return {
            "active": True,
            "rule": rule,
            "values": [],
            "counter": Counter(),
            "source_header": source_header,
            "invalid_count": len(rows),
        }

    values = []
    invalid = 0
    counter = Counter()

    for row in rows:
        raw = _string(row[source_idx] if source_idx < len(row) else "")
        digits = re.sub(r"\D", "", raw)
        if len(digits) >= 2:
            line = digits[:2]
            values.append(line)
            counter[line] += 1
        else:
            values.append("")
            invalid += 1

    return {
        "active": True,
        "rule": rule,
        "values": values,
        "counter": counter,
        "source_header": source_header,
        "invalid_count": invalid,
    }
def detect_mapping(headers: list[str]) -> dict[str, str]:
    normalized = {h: normalize_name(h) for h in headers}
    mapping: dict[str, str] = {}
    used: set[str] = set()

    # zuerst exakte Treffer
    for key, aliases in NORMALIZED_ALIASES.items():
        for header, nheader in normalized.items():
            if header in used:
                continue
            if nheader in aliases:
                mapping[key] = header
                used.add(header)
                break

    # danach vorsichtige Teiltreffer
    for key, aliases in NORMALIZED_ALIASES.items():
        if key in mapping:
            continue
        for header, nheader in normalized.items():
            if header in used or len(nheader) < 3:
                continue
            if any(
                len(alias) >= 4 and (alias in nheader or nheader in alias)
                for alias in aliases
            ):
                mapping[key] = header
                used.add(header)
                break

    return mapping


def _cell_to_json(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if value is None:
        return ""
    return value


def _string(value: Any) -> str:
    value = _cell_to_json(value)
    return str(value).strip()


def read_tabular(path: Path) -> tuple[list[str], list[list[Any]], str, str]:
    suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        if openpyxl is None:
            raise RuntimeError("Excel-Dateien benötigen das Python-Paket openpyxl.")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        first = next(rows_iter, None)
        if first is None:
            return [], [], ws.title, "Excel"

        headers = [str(v).strip() if v is not None else "" for v in first]
        rows = [list(r) for r in rows_iter]
        wb.close()
        return headers, rows, ws.title, "Excel"

    if suffix == ".csv":
        raw = path.read_bytes()
        text = None
        encoding_used = None
        for enc in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                text = raw.decode(enc)
                encoding_used = enc
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            raise RuntimeError("CSV-Zeichensatz konnte nicht erkannt werden.")

        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";"

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        all_rows = list(reader)
        if not all_rows:
            return [], [], f"CSV ({encoding_used})", "CSV"

        headers = [str(v).strip() for v in all_rows[0]]
        rows = all_rows[1:]
        return headers, rows, f"CSV · {delimiter!r} · {encoding_used}", "CSV"

    raise RuntimeError("Unterstützt werden CSV, XLSX und XLSM.")


def row_as_dict(headers: list[str], row: list[Any]) -> dict[str, Any]:
    out = {}
    for idx, header in enumerate(headers):
        out[header] = _cell_to_json(row[idx]) if idx < len(row) else ""
    return out


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    s = str(value or "").strip()
    if not s:
        return None

    # Datum ggf. vor Uhrzeit extrahieren.
    candidates = [s]
    if " " in s:
        candidates.append(s.split(" ", 1)[0])
    if "T" in s:
        candidates.append(s.split("T", 1)[0])

    for item in candidates:
        item = item.strip()
        for fmt in (
            "%Y-%m-%d",
            "%d.%m.%Y",
            "%d.%m.%y",
            "%Y/%m/%d",
            "%d/%m/%Y",
            "%m/%d/%Y",
        ):
            try:
                return datetime.strptime(item, fmt).date()
            except ValueError:
                pass
        try:
            return datetime.fromisoformat(item.replace("Z", "+00:00")).date()
        except Exception:
            pass

    return None


def get_col_index(headers: list[str], header: str | None) -> int | None:
    if not header:
        return None
    try:
        return headers.index(header)
    except ValueError:
        return None


def values_for(rows: list[list[Any]], idx: int | None):
    if idx is None:
        return []
    return [row[idx] if idx < len(row) else "" for row in rows]


def coverage(rows: list[list[Any]], idx: int | None) -> dict[str, Any]:
    total = len(rows)
    if idx is None or total == 0:
        return {"filled": 0, "total": total, "percent": 0.0}
    filled = sum(1 for r in rows if idx < len(r) and str(r[idx] if r[idx] is not None else "").strip())
    return {"filled": filled, "total": total, "percent": round(filled * 100 / total, 1)}


def get_masterdata(root: Path) -> dict[str, Any]:
    path = root / "data" / "stammdaten.json"
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception:
        return {}


def find_by_id(items: list[dict], item_id: str) -> dict | None:
    for item in items:
        if str(item.get("id", "")) == str(item_id):
            return item
    return None


def analyze_file(
    root: Path,
    source_path: Path,
    *,
    selected_bundle_id: str = "",
    selected_operator_id: str = "",
    data_type: str = "",
    delivery_period: str = "",
) -> dict[str, Any]:
    headers, rows, sheet_name, fmt = read_tabular(source_path)
    headers = [h if h else f"Spalte_{i+1}" for i, h in enumerate(headers)]
    mapping = detect_mapping(headers)




    master = get_masterdata(root)
    operator = find_by_id(master.get("verkehrsunternehmen", []), selected_operator_id)
    bundle = find_by_id(master.get("linienbuendel", []), selected_bundle_id)

    line_derivation = derive_line_values(rows, headers, mapping, operator)

    service_idx = get_col_index(headers, mapping.get("service_date"))
    dates = [d for d in (parse_date(v) for v in values_for(rows, service_idx)) if d]
    date_min = min(dates).isoformat() if dates else None
    date_max = max(dates).isoformat() if dates else None

    line_idx = get_col_index(headers, mapping.get("line"))
    line_counter = Counter(
        _string(v) for v in values_for(rows, line_idx) if _string(v)
    )

    derived_line_counter = line_derivation.get("counter", Counter())
    effective_line_counter = line_counter if line_counter else derived_line_counter

    detected_lines = [
        {"value": value, "count": count}
        for value, count in effective_line_counter.most_common(100)
    ]

    operator_idx = get_col_index(headers, mapping.get("operator"))
    vehicle_idx = get_col_index(headers, mapping.get("vehicle"))
    text_candidates = []
    for idx in (operator_idx, vehicle_idx):
        if idx is not None:
            text_candidates.extend(_string(v).lower() for v in values_for(rows[:3000], idx) if _string(v))



    checks = []

    def add(level: str, code: str, text: str):
        checks.append({"level": level, "code": code, "text": text})

    if not rows:
        add("error", "empty", "Die Datei enthält keine Datenzeilen.")

    if date_min and date_max:
        add("ok", "date_range", f"Datenzeitraum erkannt: {date_min} bis {date_max}.")
    else:
        add("warning", "date_missing", "Kein eindeutiges Datumsfeld erkannt. Der Datenzeitraum kann nicht automatisch geprüft werden.")

    if delivery_period and dates:
        months = {d.strftime("%Y-%m") for d in dates}
        if months == {delivery_period}:
            add("ok", "period_match", "Die gewählte Lieferperiode stimmt mit den erkannten Datumswerten überein.")
        else:
            sample_months = ", ".join(sorted(months)[:6])
            add(
                "warning",
                "period_mismatch",
                f"Gewählte Lieferperiode {delivery_period}; in der Datei erkannt: {sample_months}.",
            )

    if line_counter:
        preview_lines = ", ".join(x["value"] for x in detected_lines[:12])
        add("ok", "lines", f"Linien direkt aus der Datei erkannt: {preview_lines}" + (" …" if len(detected_lines) > 12 else ""))
    elif line_derivation.get("active") and detected_lines:
        preview_lines = ", ".join(x["value"] for x in detected_lines[:12])
        add(
            "ok",
            "lines_derived",
            "Linien automatisch abgeleitet: "
            + preview_lines
            + (" …" if len(detected_lines) > 12 else "")
            + ". Regel: erste zwei Stellen aus Fahrt-ID."
        )
        if line_derivation.get("invalid_count", 0):
            add(
                "warning",
                "line_derivation_invalid",
                f"{line_derivation['invalid_count']} Datensätze konnten nicht aus der Fahrt-ID einer Linie zugeordnet werden."
            )
    elif line_derivation.get("active") and not mapping.get("trip_id"):
        add(
            "warning",
            "line_derivation_no_trip",
            "Für dieses Verkehrsunternehmen ist die Linienableitung aus der Fahrt-ID vorgesehen, aber keine Fahrt-ID-Spalte wurde erkannt."
        )
    else:
        add(
            "warning",
            "line_missing",
            "Keine Linie automatisch erkannt. Für dieses Verkehrsunternehmen ist keine automatische Ableitung aus der Fahrt-ID hinterlegt."
        )

    if not mapping.get("trip_id"):
        add("warning", "trip_missing", "Keine Fahrtnummer/Fahrt-ID automatisch erkannt.")

    if operator:
        selected_name = str(operator.get("name", "")).strip()
        selected_code = str(operator.get("kuerzel", "")).strip()
        tokens = [x.lower() for x in (selected_name, selected_code) if x]
        if tokens and text_candidates:
            selected_seen = any(any(t in text for t in tokens) for text in text_candidates)
            other_seen = []
            for other in master.get("verkehrsunternehmen", []):
                if str(other.get("id", "")) == str(selected_operator_id):
                    continue
                names = [str(other.get("name", "")).strip(), str(other.get("kuerzel", "")).strip()]
                names = [x.lower() for x in names if x]
                if names and any(any(n in text for n in names) for text in text_candidates):
                    other_seen.append(str(other.get("name", "")))
            if selected_seen:
                add("ok", "operator_seen", f"Das gewählte Verkehrsunternehmen „{selected_name}“ ist in den erkannten Fahrzeug-/VU-Angaben plausibel.")
            elif other_seen:
                add("warning", "operator_other", f"In den Daten wurden Hinweise auf andere bekannte Verkehrsunternehmen gefunden: {', '.join(sorted(set(other_seen)))}.")
            else:
                add("info", "operator_unverified", f"Das gewählte Verkehrsunternehmen „{selected_name}“ konnte aus den erkannten Fahrzeug-/VU-Feldern nicht sicher bestätigt werden.")

    if bundle:
        add("info", "bundle", f"Zuordnung zum Linienbündel „{bundle.get('name', '')}“ wird als Import-Metadatum gespeichert.")

    dtype = (data_type or "").upper()
    if dtype == "ITCS":
        if not (mapping.get("scheduled_departure") or mapping.get("scheduled_arrival")):
            add("warning", "itcs_plan_time", "Keine Soll-Ankunft/Soll-Abfahrt automatisch erkannt.")
        if not (mapping.get("actual_departure") or mapping.get("actual_arrival")):
            add("warning", "itcs_actual_time", "Keine Ist-Ankunft/Ist-Abfahrt automatisch erkannt.")
    elif dtype == "AFZS":
        if not any(mapping.get(x) for x in ("boarding", "alighting", "occupancy", "load_factor")):
            add("warning", "afzs_counts", "Keine Ein-/Aussteiger-, Belegungs- oder Auslastungsfelder automatisch erkannt.")

    mapped_quality = {}
    for key, header in mapping.items():
        mapped_quality[key] = coverage(rows, get_col_index(headers, header))

    preview_rows = [
        [_cell_to_json(v) for v in row[: len(headers)]]
        for row in rows[:10]
    ]

    return {
        "ok": True,
        "filename": source_path.name,
        "format": fmt,
        "sheet_name": sheet_name,
        "row_count": len(rows),
        "column_count": len(headers),
        "headers": headers,
        "preview_rows": preview_rows,
        "mapping": mapping,
        "canonical_fields": [
            {"key": key, "label": label}
            for key, label in CANONICAL_FIELDS
        ],
        "date_range": {"from": date_min, "to": date_max},
        "detected_lines": detected_lines,
        "line_derivation": {
            "active": bool(line_derivation.get("active")),
            "rule": line_derivation.get("rule"),
            "source_header": line_derivation.get("source_header"),
            "invalid_count": int(line_derivation.get("invalid_count", 0)),
        },
        "quality": mapped_quality,
        "checks": checks,
        "selected": {
            "linienbuendel_id": selected_bundle_id,
            "verkehrsunternehmen_id": selected_operator_id,
            "datenart": data_type,
            "lieferperiode": delivery_period,
        },
    }


def required_for_controlling(data_type: str) -> list[list[str]]:
    dtype = (data_type or "").upper()
    if dtype == "ITCS":
        return [
            ["service_date"],
            ["trip_id"],
            ["stop_name", "stop_id"],
            ["scheduled_departure", "scheduled_arrival"],
            ["actual_departure", "actual_arrival"],
        ]
    if dtype == "AFZS":
        return [
            ["service_date"],
            ["stop_name", "stop_id"],
            ["boarding", "alighting", "occupancy", "load_factor"],
        ]
    return [["service_date"]]


def controlling_readiness(data_type: str, mapping: dict[str, str], *, line_available: bool = False) -> tuple[bool, list[str]]:
    missing = []
    if not line_available and "line" not in mapping:
        missing.append("Linie")
    label_by_key = dict(CANONICAL_FIELDS)
    for alternatives in required_for_controlling(data_type):
        if not any(mapping.get(key) for key in alternatives):
            missing.append(" oder ".join(label_by_key.get(k, k) for k in alternatives))
    return len(missing) == 0, missing


def commit_import(
    root: Path,
    temp_path: Path,
    *,
    original_filename: str,
    mapping: dict[str, str],
    selected_bundle_id: str,
    selected_operator_id: str,
    data_type: str,
    delivery_period: str,
) -> dict[str, Any]:
    analysis = analyze_file(
        root,
        temp_path,
        selected_bundle_id=selected_bundle_id,
        selected_operator_id=selected_operator_id,
        data_type=data_type,
        delivery_period=delivery_period,
    )

    headers, rows, sheet_name, fmt = read_tabular(temp_path)

    # Nur existierende Quellspalten akzeptieren.
    clean_mapping = {
        key: header
        for key, header in mapping.items()
        if key in dict(CANONICAL_FIELDS) and header in headers
    }

    master_for_rule = get_masterdata(root)
    selected_operator = find_by_id(master_for_rule.get("verkehrsunternehmen", []), selected_operator_id)
    line_derivation = derive_line_values(rows, headers, clean_mapping, selected_operator)
    line_available = bool(clean_mapping.get("line")) or (
        line_derivation.get("active")
        and bool(line_derivation.get("counter"))
        and int(line_derivation.get("invalid_count", 0)) == 0
    )

    ready, missing = controlling_readiness(
        data_type,
        clean_mapping,
        line_available=line_available,
    )

    now = datetime.now()
    import_id = now.strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:8]
    import_dir = root / "data" / "imports" / import_id
    import_dir.mkdir(parents=True, exist_ok=False)

    source_name = "source" + temp_path.suffix.lower()
    target_source = import_dir / source_name
    shutil.copy2(temp_path, target_source)

    indexes = {
        key: get_col_index(headers, source_header)
        for key, source_header in clean_mapping.items()
    }

    normalized_path = import_dir / "normalized.ndjson"
    with normalized_path.open("w", encoding="utf-8") as f:
        for row_index, row in enumerate(rows):
            obj = {}
            for key, idx in indexes.items():
                obj[key] = _cell_to_json(row[idx]) if idx is not None and idx < len(row) else ""

            if not obj.get("line") and line_derivation.get("active"):
                derived_values = line_derivation.get("values", [])
                if row_index < len(derived_values) and derived_values[row_index]:
                    obj["line"] = derived_values[row_index]
                    obj["_line_source"] = "derived:trip_id_first_two_digits"

            annotate_bundle(root, obj)

            f.write(json.dumps(obj, ensure_ascii=False) + "\n")

    master = get_masterdata(root)
    op = find_by_id(master.get("verkehrsunternehmen", []), selected_operator_id) or {}
    bundle = find_by_id(master.get("linienbuendel", []), selected_bundle_id) or {}

    metadata = {
        "import_id": import_id,
        "created_at": now.isoformat(timespec="seconds"),
        "original_filename": original_filename,
        "source_file": source_name,
        "normalized_file": normalized_path.name,
        "format": fmt,
        "sheet_name": sheet_name,
        "row_count": len(rows),
        "column_count": len(headers),
        "data_type": data_type,
        "delivery_period": delivery_period,
        "bundle": {"id": selected_bundle_id, "name": bundle.get("name", "")},
        "operator": {"id": selected_operator_id, "name": op.get("name", ""), "code": op.get("kuerzel", "")},
        "mapping": clean_mapping,
        "line_derivation": line_derivation.get("rule"),
        "date_range": analysis.get("date_range"),
        "detected_lines": analysis.get("detected_lines", []),
        "checks": analysis.get("checks", []),
        "controlling_ready": ready,
        "controlling_missing": missing,
        "status": "importiert" if ready else "importiert_nachbearbeitung",
    }
    (import_dir / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    history_path = root / "data" / "import_history.json"
    try:
        history = json.loads(history_path.read_text(encoding="utf-8-sig")) if history_path.exists() else []
        if not isinstance(history, list):
            history = []
    except Exception:
        history = []

    history.insert(0, {
        "import_id": import_id,
        "created_at": metadata["created_at"],
        "filename": original_filename,
        "data_type": data_type,
        "delivery_period": delivery_period,
        "bundle_name": metadata["bundle"]["name"],
        "operator_name": metadata["operator"]["name"],
        "row_count": len(rows),
        "date_range": metadata["date_range"],
        "controlling_ready": ready,
        "status": metadata["status"],
    })
    history_path.write_text(
        json.dumps(history, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    try:
        temp_path.unlink()
    except Exception:
        pass

    return {"ok": True, "metadata": metadata}


def load_history(root: Path) -> list[dict[str, Any]]:
    path = root / "data" / "import_history.json"
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        return data if isinstance(data, list) else []
    except Exception:
        return []




